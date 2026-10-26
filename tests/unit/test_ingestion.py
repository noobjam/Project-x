from docx import Document
from docx.oxml import OxmlElement

from soc_verifier.ingestion.docx import ingest_docx


def test_document_order_and_evidence_locators(tmp_path):
    path = tmp_path / "spec.docx"
    document = Document()
    document.add_heading("Reset", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("System reset ")
    paragraph.add_run("clears pending interrupts.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "RESET"
    table.cell(0, 1).text = "Active low"
    document.add_paragraph("Reset release is synchronized.")
    document.save(path)

    result = ingest_docx(path)
    assert [block.kind for block in result.blocks] == [
        "paragraph",
        "paragraph",
        "table_row",
        "paragraph",
    ]
    assert result.blocks[1].text == "System reset clears pending interrupts."
    assert result.blocks[2].text == "RESET | Active low"
    assert result.blocks[2].locator == "word/document.xml/body/*[3]/tr[1]"
    assert all(block.section == "Reset" for block in result.blocks)
    assert result.sha256 == ingest_docx(path).sha256


def test_tabs_and_breaks_are_preserved(tmp_path):
    path = tmp_path / "breaks.docx"
    document = Document()
    document.add_paragraph("pending\t1\nactive\t0")
    document.save(path)
    assert ingest_docx(path).blocks[0].text == "pending\t1\nactive\t0"


def test_content_controls_preserve_order_and_locators(tmp_path):
    path = tmp_path / "controls.docx"
    document = Document()
    document.add_heading("Reset", level=1)
    paragraph = document.add_paragraph("Reset disables DMA bus mastering.")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    document.element.body.insert(1, control)
    control.append(content)
    content.append(paragraph._p)
    document.add_paragraph("Reset clears interrupts.")
    document.save(path)
    result = ingest_docx(path)
    assert [block.text for block in result.blocks] == [
        "Reset",
        "Reset disables DMA bus mastering.",
        "Reset clears interrupts.",
    ]
    assert result.blocks[1].section == "Reset"
    assert "sdtContent" in result.blocks[1].locator
    assert len({block.id for block in result.blocks}) == 3


def test_unsupported_body_content_is_warned(tmp_path):
    path = tmp_path / "alternate.docx"
    document = Document()
    document.add_paragraph("Reset clears interrupts.")
    document.element.body.insert(1, OxmlElement("w:altChunk"))
    document.save(path)
    assert any("altChunk" in warning for warning in ingest_docx(path).extraction_warnings)

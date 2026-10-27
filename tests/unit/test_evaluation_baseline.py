import json
import subprocess
from hashlib import sha256

import pytest
from docx import Document
from openpyxl import Workbook
from pydantic import SecretStr, ValidationError

from soc_verifier.evaluation.baseline import freeze_baseline, verify_baseline
from soc_verifier.ingestion.docx import ingest_docx
from soc_verifier.settings import Settings


@pytest.fixture
def inputs(tmp_path):
    root = tmp_path / "repository"
    root.mkdir()
    (root / ".gitignore").write_text("artifacts/\nprivate/\n")
    (root / "uv.lock").write_text("version = 1\n")
    for args in [
        ["init", "-q"],
        ["add", ".gitignore", "uv.lock"],
        [
            "-c",
            "user.name=Evaluation Test",
            "-c",
            "user.email=test@example.invalid",
            "commit",
            "-qm",
            "Synthetic baseline",
        ],
    ]:
        subprocess.run(["git", "-C", str(root), *args], check=True, capture_output=True)
    private = root / "private"
    private.mkdir()
    source = private / "synthetic.docx"
    doc = Document()
    doc.add_heading("Reset", level=1)
    doc.add_paragraph("Reset clears status. Straße\tRESET\nRESET🙂")
    doc.save(source)
    reference = private / "reference.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "PRIVATE REFERENCE ANSWER"
    workbook.save(reference)
    settings = Settings(
        _env_file=None,
        model="test/model",
        provider="test/provider",
        openrouter_api_key=SecretStr("never-store-this-credential"),
    )
    return {
        "repository": root,
        "source": source,
        "reference": reference,
        "output": root / "artifacts" / "baseline",
        "settings": settings,
    }


def test_snapshot_uses_real_ingestion_and_separates_reference(inputs):
    manifest = freeze_baseline(**inputs)
    assert manifest == verify_baseline(inputs["output"], check_originals=True)
    source = inputs["output"] / manifest.runtime_input
    assert json.loads(source.read_text()) == ingest_docx(inputs["source"]).model_dump(mode="json")
    assert "PRIVATE REFERENCE ANSWER" not in source.read_text()
    assert {file.role for file in manifest.files} == {"source", "reference", "provenance"}
    assert "never-store-this-credential" not in (inputs["output"] / "manifest.json").read_text()
    assert "openrouter_api_key" not in manifest.configuration
    assert manifest.reference_status == "imported_reference_unreviewed"
    assert manifest.semantic_acceptance == "pending"
    assert manifest.historical_status == "not_supplied"


def test_dirty_repository_cannot_claim_a_frozen_commit(inputs):
    (inputs["repository"] / "uv.lock").write_text("changed\n")
    with pytest.raises(ValueError, match="Commit or stash"):
        freeze_baseline(**inputs)
    assert not inputs["output"].exists()


@pytest.mark.parametrize("destination", ["tracked-output", "../outside"])
def test_snapshot_cannot_write_private_inputs_to_publishable_paths(inputs, destination):
    inputs["output"] = inputs["repository"] / destination
    with pytest.raises(ValueError, match="ignored repository directory"):
        freeze_baseline(**inputs)
    assert not inputs["output"].exists()


def test_freeze_never_replaces_an_existing_snapshot(inputs):
    freeze_baseline(**inputs)
    original = (inputs["output"] / "manifest.json").read_bytes()
    with pytest.raises(FileExistsError):
        freeze_baseline(**inputs)
    assert (inputs["output"] / "manifest.json").read_bytes() == original


def test_original_drift_does_not_destroy_the_frozen_copy(inputs):
    manifest = freeze_baseline(**inputs)
    inputs["source"].write_bytes(b"changed original")
    assert verify_baseline(inputs["output"]) == manifest
    with pytest.raises(ValueError, match="Original input changed"):
        verify_baseline(inputs["output"], check_originals=True)


@pytest.mark.parametrize("path", ["source/specification.json", "manifest.json"])
def test_snapshot_drift_is_detected(inputs, path):
    freeze_baseline(**inputs)
    target = inputs["output"] / path
    target.write_bytes(target.read_bytes() + b" ")
    with pytest.raises(ValueError, match=r"changed|checksum"):
        verify_baseline(inputs["output"])


@pytest.mark.parametrize("change", ["version", "path", "reference_input", "config"])
def test_manifest_rejects_unsupported_or_unsafe_contracts(inputs, change):
    freeze_baseline(**inputs)
    manifest_path = inputs["output"] / "manifest.json"
    data = json.loads(manifest_path.read_bytes())
    if change == "version":
        data["format"] = "unknown-v99"
    elif change == "path":
        data["files"][0]["path"] = "source/../../outside"
    elif change == "reference_input":
        data["runtime_input"] = "reference/reference.xlsx"
    else:
        data["configuration"]["openrouter_api_key"] = "unexpected-field"
    payload = json.dumps(data).encode()
    manifest_path.write_bytes(payload)
    (inputs["output"] / "manifest.sha256").write_text(sha256(payload).hexdigest())
    with pytest.raises(ValidationError):
        verify_baseline(inputs["output"])


def test_symlink_cannot_redirect_a_snapshot_read_outside_the_bundle(inputs):
    freeze_baseline(**inputs)
    source = inputs["output"] / "source" / inputs["source"].name
    source.unlink()
    source.symlink_to(inputs["source"])
    with pytest.raises(ValueError, match="escapes"):
        verify_baseline(inputs["output"])


def history(inputs):
    root = inputs["repository"] / "private" / "history"
    root.mkdir()
    spec = ingest_docx(inputs["source"])
    (root / "manifest.json").write_text(json.dumps({"specification_sha256": spec.sha256}))
    (root / "result.json").write_text(
        json.dumps(
            {
                "plan_artifact": "plan.r000.json",
                "review_artifact": "review.r000.json",
            }
        )
    )
    (root / "specification.json").write_text(spec.model_dump_json())
    for file in ["analysis.json", "plan.r000.json", "review.r000.json"]:
        (root / file).write_text("{}")
    return root


def test_historical_outputs_are_reference_material_not_runtime_input(inputs):
    manifest = freeze_baseline(**inputs, historical_run=history(inputs))
    assert manifest.historical_status == "not_a_controlled_baseline"
    assert sum(file.role == "historical" for file in manifest.files) == 6
    assert manifest.runtime_input == "source/specification.json"


@pytest.mark.parametrize("problem", ["different_source", "escape"])
def test_incorrect_historical_provenance_is_rejected(inputs, problem):
    root = history(inputs)
    if problem == "different_source":
        (root / "manifest.json").write_text(json.dumps({"specification_sha256": "a" * 64}))
    else:
        (root / "result.json").write_text(json.dumps({"plan_artifact": "../outside.json"}))
    with pytest.raises(ValueError, match=r"different source|basenames"):
        freeze_baseline(**inputs, historical_run=root)
    assert not inputs["output"].exists()

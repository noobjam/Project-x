import pytest
from openpyxl import load_workbook

from soc_verifier.export.xlsx import export_plan


def test_export_preserves_text_ids_and_traceability(
    tmp_path, specification, analysis, verification_plan, clean_review
):
    destination = tmp_path / "plan.xlsx"
    item = verification_plan.items[0].model_copy(
        update={"id": "1.10", "objective": '=HYPERLINK("https://example.com", "text")'}
    )
    plan = verification_plan.model_copy(update={"items": [item]})
    export_plan(
        destination, specification=specification, analysis=analysis, plan=plan, review=clean_review
    )
    workbook = load_workbook(destination)
    assert workbook.sheetnames == ["Verification plan", "Requirements", "Evidence", "Findings"]
    sheet = workbook["Verification plan"]
    assert sheet["A2"].value == "1.10"
    assert sheet["D2"].data_type == "s"
    assert sheet.freeze_panes == "C2"
    assert workbook["Evidence"]["A2"].value == "P00001"
    workbook.close()
    with pytest.raises(FileExistsError):
        export_plan(
            destination,
            specification=specification,
            analysis=analysis,
            plan=plan,
            review=clean_review,
        )


def test_failed_export_leaves_no_destination_and_can_retry(
    tmp_path, monkeypatch, specification, analysis, verification_plan, clean_review
):
    destination = tmp_path / "plan.xlsx"
    arguments = {
        "specification": specification,
        "analysis": analysis,
        "plan": verification_plan,
        "review": clean_review,
    }

    def fail(self, stream):
        stream.write(b"partial workbook")
        raise OSError("simulated disk failure")

    with monkeypatch.context() as patch:
        patch.setattr("soc_verifier.export.xlsx.Workbook.save", fail)
        with pytest.raises(OSError, match="disk failure"):
            export_plan(destination, **arguments)
    assert not destination.exists()
    assert list(tmp_path.iterdir()) == []
    export_plan(destination, **arguments)
    workbook = load_workbook(destination)
    assert workbook["Verification plan"]["A2"].value == verification_plan.items[0].id
    workbook.close()


def test_source_audit_export_keeps_evidence_and_provisional_judgments_as_text(
    tmp_path, specification, analysis, verification_plan, clean_review
):
    from soc_verifier.domain.source_audit import KnowledgeApplication, SourceAudit, SourceFinding

    audit = SourceAudit(
        findings=[
            SourceFinding(
                id="F1",
                kind="lost_qualifier",
                requirement_ids=[analysis.requirements[0].id],
                evidence_ids=["P00001:U00000"],
                description="=1+1",
                recommended_action="Check controller scope",
                conflict=None,
            )
        ],
        knowledge_applications=[
            KnowledgeApplication(
                entry_id="K01",
                requirement_ids=[analysis.requirements[0].id],
                evidence_ids=["P00001:U00000"],
                assessment="uncertain",
                rationale="=1+1",
            )
        ],
    )
    path = tmp_path / "audit.xlsx"
    export_plan(
        path,
        specification=specification,
        analysis=analysis,
        plan=verification_plan,
        review=clean_review,
        source_audit=audit,
    )
    workbook = load_workbook(path)
    assert workbook["Source audit"]["E2"].value == specification.blocks[0].text
    assert workbook["Source audit"]["F2"].data_type == "s"
    assert workbook["Source audit"]["M2"].value == "agent_proposed"
    assert workbook["Knowledge application"]["E2"].data_type == "s"
    assert workbook["Knowledge application"]["F2"].value == "agent_proposed"
    workbook.close()

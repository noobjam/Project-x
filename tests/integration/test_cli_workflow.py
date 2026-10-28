import json
from unittest.mock import Mock
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pydantic import SecretStr
from typer.testing import CliRunner

from soc_verifier.cli import app
from soc_verifier.settings import Settings
from soc_verifier.storage.artifacts import ArtifactStore

pytestmark = pytest.mark.integration
runner = CliRunner()


def test_cli_plan_resume_and_export(
    tmp_path, monkeypatch, specification, analysis, verification_plan, clean_review
):
    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path / "artifacts",
        openrouter_api_key=SecretStr("test-only"),
        model="test/model",
        provider="test-endpoint",
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    agents = Mock()
    agents.configuration = settings.agent_configuration()
    agents.run.side_effect = [analysis, verification_plan, clean_review]
    monkeypatch.setattr("soc_verifier.cli.AgentRuntime", lambda settings, client, **kwargs: agents)
    source = ArtifactStore(tmp_path).write("specification.json", specification)

    generated = runner.invoke(app, ["plan", str(source)])
    assert generated.exit_code == 0, generated.output
    runs = list((settings.artifact_dir / "runs").iterdir())
    assert len(runs) == 1
    run_id = runs[0].name
    assert ArtifactStore(runs[0]).read("result.json")["status"] == "draft_ready"

    resumed = runner.invoke(app, ["resume", run_id, "--workers", "1"])
    assert resumed.exit_code == 0, resumed.output
    assert "already completed" in resumed.output
    assert agents.run.call_count == 3

    exported = runner.invoke(app, ["export", run_id, "-o", str(tmp_path / "plan.xlsx")])
    assert exported.exit_code == 0, exported.output
    assert (tmp_path / "plan.xlsx").is_file()

    changed = settings.model_copy(update={"model": "different/model"})
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: changed)
    rejected = runner.invoke(app, ["resume", run_id])
    assert rejected.exit_code == 1
    assert "configuration changed" in rejected.output
    assert agents.run.call_count == 3


def test_resume_workers_override_preserves_manifest(tmp_path, monkeypatch):
    from soc_verifier.agents.instructions import PROMPT_VERSION
    from soc_verifier.cli import _run_configuration

    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path,
        openrouter_api_key=SecretStr("test-only"),
        model="test/model",
        provider="test-endpoint",
        batch_workers=4,
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    execute = Mock()
    monkeypatch.setattr("soc_verifier.cli._execute", execute)
    run_id = str(uuid4())
    artifacts = ArtifactStore(tmp_path / "runs" / run_id)
    manifest = {"prompt_version": PROMPT_VERSION, **_run_configuration(settings)}
    artifacts.write("manifest.json", manifest)
    result = runner.invoke(app, ["resume", run_id, "--workers", "1"])
    assert result.exit_code == 0, result.output
    assert execute.call_args.args[0].batch_workers == 1
    assert execute.call_args.args[0].max_output_tokens == settings.max_output_tokens
    assert execute.call_args.args[0].max_cost_usd == settings.max_cost_usd
    assert artifacts.read("manifest.json") == manifest
    rejected = runner.invoke(app, ["resume", run_id, "--max-output-tokens", "16000"])
    assert rejected.exit_code == 1
    assert "Output-token limit changed" in rejected.output
    assert execute.call_count == 1
    lowered = runner.invoke(app, ["resume", run_id, "--max-output-tokens", "6000"])
    assert lowered.exit_code == 1 and "Output-token limit changed" in lowered.output
    assert execute.call_count == 1


def test_cli_requires_explicit_credentials(tmp_path, monkeypatch, specification):
    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path / "artifacts",
        openrouter_api_key=None,
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    source = ArtifactStore(tmp_path).write("specification.json", specification)
    result = runner.invoke(app, ["plan", str(source)])
    assert result.exit_code == 1
    assert "SOC_VERIFIER_OPENROUTER_API_KEY" in result.output
    assert not settings.artifact_dir.exists()


def test_legacy_run_cannot_resume_under_agent_engine(tmp_path, monkeypatch):
    from soc_verifier.agents.instructions import PROMPT_VERSION
    from soc_verifier.cli import _run_configuration

    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path,
        openrouter_api_key=SecretStr("test-only"),
        model="test/model",
        provider="test-endpoint",
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    execute = Mock()
    monkeypatch.setattr("soc_verifier.cli._execute", execute)
    run_id = str(uuid4())
    manifest = {"prompt_version": PROMPT_VERSION, **_run_configuration(settings)}
    del manifest["engine"]
    ArtifactStore(tmp_path / "runs" / run_id).write("manifest.json", manifest)
    result = runner.invoke(app, ["resume", run_id])
    assert result.exit_code == 1 and "configuration changed" in result.output
    execute.assert_not_called()


def test_export_uses_final_analysis_revision(
    tmp_path, monkeypatch, specification, analysis, verification_plan, clean_review
):
    settings = Settings(_env_file=None, artifact_dir=tmp_path / "artifacts")
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    run_id = str(uuid4())
    artifacts = ArtifactStore(settings.artifact_dir / "runs" / run_id)
    revised = analysis.model_copy(
        update={
            "requirements": [
                *analysis.requirements,
                analysis.requirements[0].model_copy(update={"id": "REQ-RST-002"}),
            ]
        }
    )
    plan = verification_plan.model_copy(
        update={
            "items": [
                verification_plan.items[0].model_copy(
                    update={"requirement_ids": ["REQ-RST-001", "REQ-RST-002"]}
                )
            ]
        }
    )
    for name, artifact in [
        ("specification.json", specification),
        ("analysis.json", analysis),
        ("analysis.r001.json", revised),
        ("plan.r001.json", plan),
        ("review.r001.json", clean_review),
    ]:
        artifacts.write(name, artifact)
    artifacts.write(
        "result.json",
        {
            "analysis_artifact": "analysis.r001.json",
            "plan_artifact": "plan.r001.json",
            "review_artifact": "review.r001.json",
        },
    )
    destination = tmp_path / "revised.xlsx"
    result = runner.invoke(app, ["export", run_id, "-o", str(destination)])
    assert result.exit_code == 0, result.output
    workbook = load_workbook(destination)
    assert workbook["Run metadata"]["B2"].value == "unavailable (legacy run)"
    assert workbook["Requirements"].max_row == 3
    assert workbook["Requirements"]["A3"].value == "REQ-RST-002"
    workbook.close()


def test_unknown_result_version_cannot_export(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, artifact_dir=tmp_path)
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    run_id = str(uuid4())
    artifacts = ArtifactStore(tmp_path / "runs" / run_id)
    artifacts.write("result.json", {"schema_version": "999"})
    result = runner.invoke(app, ["export", run_id, "-o", str(tmp_path / "out.xlsx")])
    assert result.exit_code == 1 and "Unknown run result schema version" in result.output
    assert not (tmp_path / "out.xlsx").exists()


def test_cli_freezes_knowledge_and_exports_retrieval(
    tmp_path,
    monkeypatch,
    specification,
    analysis,
    verification_plan,
    clean_review,
    knowledge_catalog,
):
    from soc_verifier.agents.tasks import ArtifactReference

    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path / "artifacts",
        openrouter_api_key=SecretStr("test-only"),
        model="test/model",
        provider="test-endpoint",
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    agents = Mock()
    agents.configuration = settings.agent_configuration()
    agents.run.side_effect = [analysis, verification_plan, clean_review]
    monkeypatch.setattr("soc_verifier.cli.AgentRuntime", lambda *args, **kwargs: agents)
    source = ArtifactStore(tmp_path).write("specification.json", specification)
    arguments = ["plan", str(source), "--knowledge-catalog", str(knowledge_catalog)]
    rejected = runner.invoke(app, arguments)
    assert rejected.exit_code == 1 and "No approved knowledge" in rejected.output
    agents.run.assert_not_called()
    generated = runner.invoke(app, [*arguments, "--allow-proposed-knowledge"])
    assert generated.exit_code == 0, generated.output
    root = next((settings.artifact_dir / "runs").iterdir())
    store = ArtifactStore(root)
    ref = ArtifactReference.capture(store, "knowledge.json")
    assert store.read("manifest.json")["knowledge"] == ref.model_dump()
    for call in agents.run.call_args_list:
        assert call.kwargs["task"].inputs.evidence["knowledge"] == ref
    # Authoring edits must not affect an existing run's frozen snapshot.
    knowledge_catalog.write_text("{}")
    resumed = runner.invoke(app, ["resume", root.name])
    assert resumed.exit_code == 0, resumed.output
    assert agents.run.call_count == 3
    access = {
        "format": "knowledge-access-v1",
        "created_at": "2026-10-28T00:00:00Z",
        "role": "planner",
        "task_id": "synthetic-task",
        "tool": "get_knowledge_entries",
        "entry_ids": ["K01"],
        "snapshot_sha256": ref.sha256,
        "offset": 0,
    }
    store.write("knowledge-access.synthetic.json", access)
    reported = runner.invoke(app, ["knowledge", root.name])
    assert reported.exit_code == 0 and json.loads(reported.output) == [access]
    destination = tmp_path / "knowledge-plan.xlsx"
    exported = runner.invoke(app, ["export", root.name, "-o", str(destination)])
    assert exported.exit_code == 0, exported.output
    workbook = load_workbook(destination)
    assert workbook["Knowledge retrieval"]["D2"].value == "K01"
    assert workbook["Knowledge retrieval"]["E2"].value == ref.sha256
    workbook.close()
    store.path("knowledge.json").unlink()
    rejected = runner.invoke(app, ["resume", root.name])
    assert rejected.exit_code == 1 and "knowledge snapshot changed" in rejected.output
    assert agents.run.call_count == 3


def test_cli_proposed_flag_requires_catalog(tmp_path, specification):
    source = ArtifactStore(tmp_path).write("specification.json", specification)
    result = runner.invoke(app, ["plan", str(source), "--allow-proposed-knowledge"])
    assert result.exit_code == 1 and "requires --knowledge-catalog" in result.output

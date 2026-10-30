import json
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from soc_verifier.cli import app
from soc_verifier.domain.source_audit import SourceAudit
from soc_verifier.settings import Settings
from soc_verifier.storage.artifacts import ArtifactStore
from soc_verifier.workflows.spec_to_plan.accounting import read_accounting
from soc_verifier.workflows.spec_to_plan.workflow import PlanningWorkflow

pytestmark = pytest.mark.integration


@pytest.fixture
def accounting_draft(passage_accounting_draft):
    return passage_accounting_draft


class Agents:
    def __init__(self, *responses):
        self.configuration = {}
        self.responses = iter(responses)
        self.calls = []

    def run(self, **kwargs):
        self.calls.append(kwargs)
        if issubclass(kwargs["schema"], SourceAudit):
            return SourceAudit(findings=[], knowledge_applications=[])
        response = next(self.responses)
        if isinstance(response, Exception):
            raise response
        return kwargs["schema"].model_validate(response.model_dump())


def test_accounting_is_bound_before_planning_and_reused(
    tmp_path, specification, accounting_draft, verification_plan, clean_review
):
    artifacts = ArtifactStore(tmp_path)
    first = Agents(accounting_draft, RuntimeError("interrupted"))
    with pytest.raises(RuntimeError, match="interrupted"):
        PlanningWorkflow(first, artifacts, max_revisions=0, source_accounting=True).run(
            specification
        )
    report = read_accounting(artifacts, "source-accounting.r000.json")
    assert report.semantic_acceptance == "unreviewed"
    assert (
        "Conditions:"
        in json.loads(first.calls[1]["context"])["analysis"]["requirements"][0]["statement"]
    )
    assert "source_accounting" in first.calls[1]["task"].inputs.evidence
    second = Agents(verification_plan, clean_review)
    PlanningWorkflow(second, artifacts, max_revisions=0, source_accounting=True).run()
    assert [c["role"] for c in second.calls] == ["planner", "reviewer"]
    assert (
        artifacts.read("result.json")["source_accounting_artifact"] == "source-accounting.r000.json"
    )
    with pytest.raises(ValueError, match="changed"):
        PlanningWorkflow(Agents(), artifacts, max_revisions=0, source_accounting=False).run()


@pytest.mark.parametrize(
    "target",
    ["source-accounting.r000.json", "analysis.json", "obligations.r000.json", "specification.json"],
)
def test_tampering_rejected_without_inference(
    tmp_path, specification, accounting_draft, verification_plan, clean_review, target
):
    artifacts = ArtifactStore(tmp_path)
    PlanningWorkflow(
        Agents(accounting_draft, verification_plan, clean_review),
        artifacts,
        max_revisions=0,
        source_accounting=True,
    ).run(specification)
    data = artifacts.read(target)
    if target.startswith("source-accounting"):
        data["obligations"][0]["evidence"][0]["start"] = 1
    elif target == "analysis.json":
        data["requirements"][0]["statement"] += " invented"
    elif target.startswith("obligations"):
        data["obligations"][0]["exceptions"] = ["invented"]
    else:
        data["blocks"][0]["text"] += " changed"
    artifacts.path(target).write_text(json.dumps(data))
    with pytest.raises(ValueError):
        read_accounting(artifacts, "source-accounting.r000.json")


def test_invalid_source_unit_stops_before_planner(tmp_path, specification, accounting_draft):
    invalid = accounting_draft.model_copy(
        update={
            "obligations": [
                accounting_draft.obligations[0].model_copy(update={"evidence_ids": ["made-up"]})
            ]
        }
    )
    agents = Agents(invalid)
    artifacts = ArtifactStore(tmp_path)
    with pytest.raises(ValueError):
        PlanningWorkflow(agents, artifacts, max_revisions=0, source_accounting=True).run(
            specification
        )
    assert len(agents.calls) == 1
    assert not artifacts.exists("analysis.json")
    assert not artifacts.exists("obligations.r000.completed.json")


def test_large_source_rejected_before_inference(tmp_path, specification):
    blocks = [
        specification.blocks[0].model_copy(update={"id": f"P{i}", "text": "x" * 800})
        for i in range(4)
    ]
    agents = Agents()
    with pytest.raises(ValueError, match="one bounded"):
        PlanningWorkflow(
            agents,
            ArtifactStore(tmp_path),
            max_revisions=0,
            chunk_max_characters=1000,
            source_accounting=True,
        ).run(specification.model_copy(update={"blocks": blocks}))
    assert not agents.calls


def test_unaccounted_source_prevents_clean_status(
    tmp_path, specification, accounting_draft, verification_plan, clean_review
):
    other = specification.blocks[0].model_copy(
        update={"id": "P00002", "text": "State is retained."}
    )
    spec = specification.model_copy(update={"blocks": [*specification.blocks, other]})
    result = PlanningWorkflow(
        Agents(accounting_draft, verification_plan, clean_review),
        ArtifactStore(tmp_path),
        max_revisions=0,
        source_accounting=True,
    ).run(spec)
    assert result["status"] == "needs_review"


def test_offline_cli_inspection_export_and_legacy_unavailable(
    tmp_path, monkeypatch, specification, accounting_draft, verification_plan, clean_review
):
    settings = Settings(_env_file=None, artifact_dir=tmp_path)
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    run_id = str(uuid4())
    artifacts = ArtifactStore(tmp_path / "runs" / run_id)
    PlanningWorkflow(
        Agents(accounting_draft, verification_plan, clean_review),
        artifacts,
        max_revisions=0,
        source_accounting=True,
    ).run(specification)
    runner = CliRunner()
    inspected = runner.invoke(app, ["accounting", run_id])
    assert inspected.exit_code == 0, inspected.output
    assert json.loads(inspected.output)["semantic_acceptance"] == "unreviewed"
    audited = runner.invoke(app, ["audit", run_id])
    assert audited.exit_code == 0, audited.output
    assert json.loads(audited.output)["judgment_state"] == "agent_proposed"
    exported = runner.invoke(app, ["export", run_id, "-o", str(tmp_path / "plan.xlsx")])
    assert exported.exit_code == 0, exported.output
    with (tmp_path / "plan.xlsx").open("rb") as stream:
        workbook = load_workbook(stream)
        assert workbook["Source accounting"]["F2"].value == specification.blocks[0].text
        assert workbook["Source accounting"]["G2"].data_type == "s"
        assert "Source audit" in workbook.sheetnames
        workbook.close()
    unavailable = runner.invoke(app, ["accounting", run_id, "--revision", "1"])
    assert unavailable.exit_code == 0 and "unavailable" in unavailable.output
    unavailable = runner.invoke(app, ["audit", run_id, "--revision", "1"])
    assert unavailable.exit_code == 0 and "unavailable" in unavailable.output


def test_cli_opt_in_is_persisted_and_required_on_resume(
    tmp_path, monkeypatch, specification, accounting_draft, verification_plan, clean_review
):
    from pydantic import SecretStr

    settings = Settings(
        _env_file=None,
        artifact_dir=tmp_path / "artifacts",
        openrouter_api_key=SecretStr("test-only"),
        model="test/model",
        provider="test-endpoint",
    )
    monkeypatch.setattr("soc_verifier.cli.Settings", lambda: settings)
    agents = Agents(accounting_draft, verification_plan, clean_review)

    def runtime(config, client, **kwargs):
        agents.configuration = config.agent_configuration()
        return agents

    monkeypatch.setattr("soc_verifier.cli.AgentRuntime", runtime)
    source = ArtifactStore(tmp_path).write("source.json", specification)
    runner = CliRunner()
    result = runner.invoke(app, ["plan", str(source), "--source-accounting", "--analysis-only"])
    assert result.exit_code == 0, result.output
    assert "analysis_ready" in result.output
    run = next((settings.artifact_dir / "runs").iterdir())
    assert not ArtifactStore(run).exists("result.json")
    assert [c["role"] for c in agents.calls] == ["analyst", "reviewer"]
    audited = runner.invoke(app, ["audit", run.name])
    assert audited.exit_code == 0, audited.output
    assert ArtifactStore(run).read("manifest.json")["source_accounting"] is True
    denied = runner.invoke(app, ["resume", run.name])
    assert denied.exit_code == 1 and "configuration changed" in denied.output
    resumed = runner.invoke(app, ["resume", run.name, "--source-accounting"])
    assert resumed.exit_code == 0, resumed.output
    assert len(agents.calls) == 4


def test_revised_analysis_keeps_matching_accounting_and_originals(
    tmp_path, specification, accounting_draft, verification_plan, clean_review
):
    from soc_verifier.domain.models import PlanReview, ReviewFinding

    review = PlanReview(
        findings=[
            ReviewFinding(
                severity="error", target="analysis", item_ids=[], description="Check extraction"
            )
        ]
    )
    updated = accounting_draft.model_copy(update={"gaps": ["Timing unresolved"]})
    artifacts = ArtifactStore(tmp_path)
    agents = Agents(
        accounting_draft, verification_plan, review, updated, verification_plan, clean_review
    )
    PlanningWorkflow(agents, artifacts, max_revisions=1, source_accounting=True).run(specification)
    result = artifacts.read("result.json")
    assert result["analysis_artifact"] == "analysis.r001.json"
    assert result["source_accounting_artifact"] == "source-accounting.r001.json"
    assert result["source_audit_artifact"] == "source-audit.r001.json"
    audit_inputs = artifacts.read("source-audit.r001.task.json")["inputs"]["evidence"]
    assert audit_inputs["analysis"]["name"] == "analysis.r001.json"
    assert audit_inputs["source_accounting"]["name"] == "source-accounting.r001.json"
    old = read_accounting(artifacts, "source-accounting.r000.json")
    new = read_accounting(artifacts, "source-accounting.r001.json")
    assert old.analysis_sha256 != new.analysis_sha256
    assert artifacts.read("analysis.json")["gaps"] == []
    assert artifacts.read("analysis.r001.json")["gaps"] == ["Timing unresolved"]
    PlanningWorkflow(Agents(), artifacts, max_revisions=1, source_accounting=True).run()

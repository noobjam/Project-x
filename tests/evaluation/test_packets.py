import json
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pydantic import ValidationError

from soc_verifier.evaluation.baseline import freeze_baseline
from soc_verifier.evaluation.packets import (
    PrivateSelection,
    PublicRecipe,
    fetch_public,
    load_packet,
    prepare_private,
    prepare_public,
)
from soc_verifier.ingestion.docx import ingest_docx
from soc_verifier.settings import Settings


@pytest.fixture
def private_case(evaluation_repo):
    private = evaluation_repo / "artifacts"
    private.mkdir()
    source = private / "source.docx"
    document = Document()
    document.add_heading("Reset", level=1)
    document.add_paragraph("Reset clears status. Straße🙂")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Signal", "Effect"
    table.cell(1, 0).text, table.cell(1, 1).text = "RESET", "Clear status"
    document.save(source)
    reference = private / "reference.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "REFERENCE ANSWER MUST NOT ENTER SOURCE"
    workbook.save(reference)
    baseline = private / "baseline"
    freeze_baseline(
        repository=evaluation_repo,
        source=source,
        reference=reference,
        output=baseline,
        settings=Settings(_env_file=None, model="synthetic/model", provider="synthetic/provider"),
    )
    specification = ingest_docx(source)
    selection = PrivateSelection(
        baseline_manifest_sha256=sha256((baseline / "manifest.json").read_bytes()).hexdigest(),
        cases=[
            {
                "case_id": "reset-table",
                "scope": "Synthetic reset table",
                "block_ids": [block.id for block in specification.blocks],
                "limitations": ["Synthetic fixture, not engineering acceptance"],
            }
        ],
    )
    return baseline, selection, specification


def test_private_packet_preserves_actual_docx_rows_and_separates_answers(
    evaluation_repo, private_case
):
    baseline, selection, specification = private_case
    (path,) = prepare_private(baseline, selection, evaluation_repo / "artifacts" / "packets")
    manifest, projected = load_packet(path)
    assert projected == specification
    assert any(block.kind == "table_row" for block in projected.blocks)
    assert set(manifest.block_documents.values()) == {specification.sha256}
    assert manifest.context_state == "proposed"
    assert list((path / "source").iterdir()) == [path / manifest.runtime_input]
    assert "REFERENCE ANSWER" not in (path / manifest.runtime_input).read_text()


@pytest.mark.parametrize("problem", ["baseline", "unknown_id", "duplicate_id", "publishable"])
def test_private_selection_rejects_wrong_provenance_and_unsafe_output(
    evaluation_repo, private_case, problem
):
    baseline, selection, _ = private_case
    data = selection.model_dump()
    if problem == "baseline":
        data["baseline_manifest_sha256"] = "0" * 64
    elif problem == "unknown_id":
        data["cases"][0]["block_ids"].append("UNKNOWN")
    elif problem == "duplicate_id":
        data["cases"][0]["block_ids"].append(data["cases"][0]["block_ids"][0])
    output = evaluation_repo / ("publishable" if problem == "publishable" else "artifacts/packets")
    with pytest.raises(ValueError):
        prepare_private(baseline, PrivateSelection.model_validate(data), output)


def public_fixture(root):
    files = {
        "LICENSE": b"Synthetic fixture license\n",
        "a.md": b"# Reset\r\nReset clears status.\r\n",
        "b.md": "# Wake\nStraße\tWAKE🙂\n".encode(),
    }
    for name, data in files.items():
        (root / name).write_bytes(data)
    recipe = PublicRecipe(
        repository="https://github.com/lowRISC/opentitan",
        revision="a" * 40,
        files={name: sha256(data).hexdigest() for name, data in files.items()},
        cases=[
            {
                "case_id": "integration",
                "scope": "Synthetic reset and wake",
                "selections": [{"path": name, "first": 1, "last": 2} for name in ("a.md", "b.md")],
                "limitations": [],
            }
        ],
    )
    return recipe


def test_public_projection_keeps_document_identity_and_exact_line_text(evaluation_repo):
    source = evaluation_repo / "artifacts" / "public"
    source.mkdir(parents=True)
    recipe = public_fixture(source)
    (packet,) = prepare_public(recipe, source, evaluation_repo / "artifacts" / "packets")
    manifest, specification = load_packet(packet)
    assert len(specification.blocks) == 4
    assert len(set(manifest.block_documents.values())) == 2
    assert specification.blocks[-1].text == "Straße\tWAKE🙂"
    assert specification.blocks[-1].locator == "b.md:L2"
    assert manifest.sources[-1].revision == "a" * 40
    with pytest.raises(FileExistsError):
        prepare_public(recipe, source, packet.parent)


@pytest.mark.parametrize("problem", ["hash", "range", "overlap", "path", "version"])
def test_public_fixture_rejects_drift_and_invalid_ranges(evaluation_repo, problem):
    source = evaluation_repo / "artifacts" / "public"
    source.mkdir(parents=True)
    recipe = public_fixture(source)
    data = recipe.model_dump()
    if problem == "hash":
        (source / "a.md").write_text("changed")
    elif problem == "range":
        data["cases"][0]["selections"][0]["last"] = 999
    elif problem == "overlap":
        data["cases"][0]["selections"].append(data["cases"][0]["selections"][0])
    elif problem == "path":
        data["files"]["../outside"] = "a" * 64
    else:
        data["format"] = "unknown-v2"
    with pytest.raises((ValueError, ValidationError)):
        prepare_public(
            PublicRecipe.model_validate(data), source, evaluation_repo / "artifacts" / "packets"
        )


def test_cached_public_download_is_offline_and_rejects_changed_cache(evaluation_repo):
    source = evaluation_repo / "artifacts" / "public"
    source.mkdir(parents=True)
    recipe = public_fixture(source)
    fetch_public(recipe, source)
    (source / "a.md").write_text("changed")
    with pytest.raises(ValueError, match="hash mismatch"):
        fetch_public(recipe, source)


def test_packet_manifest_and_source_drift_fail_explicitly(evaluation_repo):
    source = evaluation_repo / "artifacts" / "public"
    source.mkdir(parents=True)
    (packet,) = prepare_public(
        public_fixture(source), source, evaluation_repo / "artifacts" / "packets"
    )
    path = packet / "source/specification.json"
    path.write_text(path.read_text() + " ")
    with pytest.raises(ValueError, match="source hash mismatch"):
        load_packet(packet)
    (packet / "packet.json").write_text(json.dumps({"format": "unknown"}))
    with pytest.raises(ValueError, match="checksum"):
        load_packet(packet)


def test_public_fetch_uses_pinned_urls_and_checks_bytes_before_writing(
    evaluation_repo, monkeypatch
):
    source = evaluation_repo / "artifacts" / "public"
    source.mkdir(parents=True)
    recipe = public_fixture(source)
    expected = {name: (source / name).read_bytes() for name in recipe.files}
    urls = []

    def download(url, timeout):
        urls.append(url)
        prefix = f"https://raw.githubusercontent.com/lowRISC/opentitan/{recipe.revision}/"
        assert url.startswith(prefix) and timeout == 30
        return BytesIO(expected[url.removeprefix(prefix)])

    monkeypatch.setattr("urllib.request.urlopen", download)
    destination = source.parent / "downloaded"
    fetch_public(recipe, destination)
    assert len(urls) == len(recipe.files)
    assert (destination / "b.md").read_bytes() == expected["b.md"]
    monkeypatch.setattr("urllib.request.urlopen", lambda *args, **kwargs: BytesIO(b"wrong bytes"))
    rejected = source.parent / "bad-download"
    with pytest.raises(ValueError, match="hash mismatch"):
        fetch_public(recipe, rejected)
    assert not (rejected / "LICENSE").exists()


def test_checked_in_public_recipe_is_valid_and_pins_integration_dependencies():
    path = Path(__file__).parents[2] / "evals" / "opentitan.json"
    recipe = PublicRecipe.model_validate_json(path.read_bytes())
    integration = next(case for case in recipe.cases if case.case_id == "opentitan-usb-wake")
    paths = {selection.path for selection in integration.selections}
    assert any("/usbdev/" in path for path in paths)
    assert any("/pinmux/" in path for path in paths)
    assert any("/pwrmgr/" in path for path in paths)

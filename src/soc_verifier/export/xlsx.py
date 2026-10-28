"""Excel review views; exported text is never interpreted as a formula."""

import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from soc_verifier.domain.accounting import SourceAccounting
from soc_verifier.domain.models import (
    PlanReview,
    SpecAnalysis,
    Specification,
    VerificationPlan,
    validate_sources,
    validate_traceability,
)


def export_plan(
    path: Path,
    *,
    specification: Specification,
    analysis: SpecAnalysis,
    plan: VerificationPlan,
    review: PlanReview,
    task_metadata: str | None = None,
    knowledge_accesses: list[dict[str, Any]] | None = None,
    source_accounting: SourceAccounting | None = None,
) -> None:
    validate_sources(analysis, specification)
    validate_traceability(plan, analysis)
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "Verification plan"
    _row(
        sheet,
        [
            "ID",
            "Requirement IDs",
            "Feature",
            "Objective",
            "Method",
            "Preconditions",
            "Stimulus",
            "Expected results",
            "Coverage intent",
            "Environment needs",
            "Priority",
            "Open questions",
            "Status",
        ],
    )
    for item in plan.items:
        _row(
            sheet,
            [
                item.id,
                "\n".join(item.requirement_ids),
                item.feature,
                item.objective,
                item.method,
                "\n".join(item.preconditions),
                "\n".join(item.stimulus),
                "\n".join(item.expected_results),
                "\n".join(item.coverage_intent),
                "\n".join(item.environment_needs),
                item.priority,
                "\n".join(item.open_questions),
                "Draft — requires human review",
            ],
        )
    requirements = workbook.create_sheet("Requirements")
    _row(requirements, ["ID", "Feature", "Statement", "Basis", "Source IDs", "Questions"])
    for requirement in analysis.requirements:
        _row(
            requirements,
            [
                requirement.id,
                requirement.feature,
                requirement.statement,
                requirement.basis,
                "\n".join(requirement.source_ids),
                "\n".join(requirement.open_questions),
            ],
        )
    evidence = workbook.create_sheet("Evidence")
    _row(evidence, ["Source ID", "Document", "SHA-256", "Section", "Locator", "Text"])
    used_sources = {source for item in analysis.requirements for source in item.source_ids}
    for block in specification.blocks:
        if block.id in used_sources:
            _row(
                evidence,
                [
                    block.id,
                    specification.document_name,
                    specification.sha256,
                    block.section,
                    block.locator,
                    block.text,
                ],
            )
    findings = workbook.create_sheet("Findings")
    _row(findings, ["Origin", "Severity", "Item IDs", "Description"])
    for warning in specification.extraction_warnings:
        _row(findings, ["extraction", "warning", "", warning])
    for gap in analysis.gaps:
        _row(findings, ["specification", "warning", "", gap])
    for finding in review.findings:
        _row(
            findings,
            [
                f"review/{finding.target}",
                finding.severity,
                "\n".join(finding.item_ids),
                finding.description,
            ],
        )

    if task_metadata is not None:
        metadata = workbook.create_sheet("Run metadata")
        _row(metadata, ["Field", "Availability"])
        _row(metadata, ["Task input and spending attribution", task_metadata])
        _row(metadata, ["Source accounting", "unreviewed" if source_accounting else "unavailable"])

    if source_accounting is not None:
        sheet = workbook.create_sheet("Source accounting")
        _row(
            sheet,
            [
                "Record",
                "Requirement",
                "Source ID",
                "Start",
                "End",
                "Source text",
                "Conditions",
                "Constraints",
                "Exceptions",
                "Reason",
            ],
        )
        blocks = {block.id: block for block in specification.blocks}
        for obligation in source_accounting.obligations:
            for span in obligation.evidence:
                _row(
                    sheet,
                    [
                        "Unreviewed evidence",
                        obligation.id,
                        span.block_id,
                        str(span.start),
                        str(span.end),
                        blocks[span.block_id].text[span.start : span.end],
                        "\n".join(obligation.conditions),
                        "\n".join(obligation.constraints),
                        "\n".join(obligation.exceptions),
                        "",
                    ],
                )
        for disposition in source_accounting.dispositions:
            for identifier in disposition.source_ids:
                _row(
                    sheet,
                    [
                        "Proposed " + disposition.status,
                        "",
                        identifier,
                        "",
                        "",
                        blocks[identifier].text,
                        "",
                        "",
                        "",
                        disposition.reason,
                    ],
                )
        for span in source_accounting.unaccounted:
            _row(
                sheet,
                [
                    "Unaccounted",
                    "",
                    span.block_id,
                    str(span.start),
                    str(span.end),
                    blocks[span.block_id].text[span.start : span.end],
                    "",
                    "",
                    "",
                    "No quotation or disposition accounts for this text.",
                ],
            )
        _row(sheet, ["Quotations establish source location, not complete or correct extraction."])

    if knowledge_accesses:
        knowledge = workbook.create_sheet("Knowledge retrieval")
        _row(knowledge, ["Role", "Task", "Action", "Entry IDs", "Snapshot SHA-256", "Offset"])
        for access in knowledge_accesses:
            _row(
                knowledge,
                [
                    access["role"],
                    access["task_id"],
                    access["tool"],
                    "\n".join(access["entry_ids"]),
                    access["snapshot_sha256"],
                    str(access["offset"]),
                ],
            )
        _row(knowledge, ["Retrieved guidance is not proof of correct application."])

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "C2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column_index, column in enumerate(worksheet.columns, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = 30
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Build beside the destination, then publish atomically without replacing it.
    # An interrupted save must not leave a partial workbook at the requested path.
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as file:
            temporary = Path(file.name)
            workbook.save(file)
            file.flush()
            os.fsync(file.fileno())
        os.link(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        workbook.close()


def _row(sheet: Worksheet, values: list[str]) -> None:
    if any(len(value) > 32_767 for value in values):
        raise ValueError("Excel cell length exceeded; the full text remains in JSON")
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.data_type = "s"
